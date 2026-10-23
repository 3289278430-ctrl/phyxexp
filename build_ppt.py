# -*- coding: utf-8 -*-
"""生成 phyxexp 参赛 PPT（初版，10 页）。"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

PRIMARY = RGBColor(0x1F, 0x38, 0x64)
ACCENT  = RGBColor(0xF5, 0x85, 0x18)
GRAY    = RGBColor(0x40, 0x40, 0x40)
LIGHT   = RGBColor(0xED, 0xF1, 0xF7)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
GREEN   = RGBColor(0x2E, 0x7D, 0x32)
RED     = RGBColor(0xC0, 0x39, 0x2B)
BLUE2   = RGBColor(0x4C, 0x78, 0xA8)
FONT = "Microsoft YaHei"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def style(run, size=18, bold=False, color=GRAY, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn('a:ea'))
    if ea is None:
        ea = rPr.makeelement(qn('a:ea'), {})
        rPr.append(ea)
    ea.set('typeface', name)

def add_text(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        p.space_after = Pt(space_after)
        if isinstance(ln, str):
            ln = (ln, 18, False, GRAY)
        text, size, bold, color = ln
        r = p.add_run(); r.text = text
        style(r, size, bold, color)
    return tb

def add_title(slide, text, sub=None):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    band.fill.solid(); band.fill.fore_color.rgb = PRIMARY
    band.line.fill.background()
    band.shadow.inherit = False
    tf = band.text_frame
    tf.margin_left = Inches(0.6); tf.margin_top = Inches(0.12); tf.margin_bottom = Inches(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = text
    style(r, 26, True, WHITE)
    if sub:
        p2 = tf.add_paragraph()
        r2 = p2.add_run(); r2.text = sub
        style(r2, 12, False, RGBColor(0xD5, 0xDE, 0xEC))
    ln = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.25), Inches(2.4), Inches(0.07))
    ln.fill.solid(); ln.fill.fore_color.rgb = ACCENT
    ln.line.fill.background(); ln.shadow.inherit = False

def add_card(slide, x, y, w, h, header, body_lines, header_color=PRIMARY):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid(); card.fill.fore_color.rgb = LIGHT
    card.line.color.rgb = RGBColor(0xC9, 0xD4, 0xE4); card.line.width = Pt(1)
    card.shadow.inherit = False
    tf = card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.18); tf.margin_right = Inches(0.18)
    tf.margin_top = Inches(0.12); tf.margin_bottom = Inches(0.1)
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = header
    style(r, 17, True, header_color)
    p.space_after = Pt(4)
    for ln in body_lines:
        p2 = tf.add_paragraph()
        p2.space_after = Pt(3)
        r2 = p2.add_run(); r2.text = ln
        style(r2, 12.5, False, GRAY)
    return card

# ============ Slide 1 封面 ============
s = prs.slides.add_slide(BLANK)
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg.fill.solid(); bg.fill.fore_color.rgb = PRIMARY; bg.line.fill.background(); bg.shadow.inherit = False
add_text(s, 0.9, 2.0, 11.5, 1.4, [("phyxexp", 66, True, WHITE)])
add_text(s, 0.9, 3.4, 11.5, 0.8, [("受约束的学生实验数据处理 Agent / Skill", 26, True, ACCENT)])
add_text(s, 0.9, 4.4, 11.5, 1.6, [
    ("第四届开放原子大赛 · 图1赛道：开源基础软件与解决方案", 16, False, RGBColor(0xD5,0xDE,0xEC)),
    ("约束，本身就是可信度的来源", 18, True, WHITE),
], space_after=10)

# ============ Slide 2 痛点 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "痛点：AI 做数据分析的三大病")
add_card(s, 0.7, 1.8, 3.8, 3.6, "① 会编数字", [
    "张口就给一个点估计，",
    "数值无来源、无推导，",
    "更没有不确定度。",
    "",
    "你无法判断它是对是错。",
], RED)
add_card(s, 5.0, 1.8, 3.8, 3.6, "② 结论不可复现", [
    "同一个问题换种问法，",
    "答案可能就变了。",
    "",
    "没有运行记录，",
    "无法回查、无法审计。",
], ACCENT)
add_card(s, 9.3, 1.8, 3.8, 3.6, "③ 缺物理一致性检查", [
    "不看残差、不查截距、",
    "不管量纲自洽。",
    "",
    "一个「看起来合理」的数，",
    "可能违反物理规律。",
], RED)
add_text(s, 0.7, 5.8, 12.0, 0.9, [
    ("核心问题：大模型「自由写代码算数」，把「表达流畅」错当成「计算可信」。", 16, True, PRIMARY),
])

# ============ Slide 3 创意方案 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "创意方案：受约束 Agent 编排层")
add_text(s, 0.7, 1.7, 12.0, 1.0, [
    ("不是让大模型更会算数，而是——", 20, False, GRAY),
    ("不让它算数。", 30, True, PRIMARY),
], space_after=4)
add_card(s, 0.7, 3.0, 3.8, 2.9, "只挑白名单工具", [
    "计算能力只能来自",
    "审查过的领域工具链",
    "（phyxexp / TOOL_REGISTRY），",
    "其余一律拒绝。",
])
add_card(s, 5.0, 3.0, 3.8, 2.9, "执行可复现计算", [
    "不确定度、R²、残差图，",
    "每一步都有 run_id +",
    "输入哈希，重放逐位一致。",
])
add_card(s, 9.3, 3.0, 3.8, 2.9, "结论必须基于数据", [
    "结论里的每个数字都要",
    "能溯源到工具输出",
    "（grounding 校验），",
    "编不进去。",
])
add_text(s, 0.7, 6.15, 12.0, 0.7, [("约束 = 可信度：把「会编数」的自由，换成「可溯源」的纪律。", 16, True, ACCENT)])

# ============ Slide 4 架构图 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "系统架构：三层受约束")
def layer(y, header, body, color=BLUE2):
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(y), Inches(10.1), Inches(1.05))
    box.fill.solid(); box.fill.fore_color.rgb = LIGHT
    box.line.color.rgb = color; box.line.width = Pt(1.5)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = Inches(0.25); tf.margin_top = Inches(0.06); tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = header + "  "
    style(r, 16, True, color)
    r2 = p.add_run(); r2.text = body
    style(r2, 13, False, GRAY)
layer(1.75, "宿主大模型", "Claude Code（SKILL.md）/ ChatGPT Custom GPT（Actions）—— 只做「选实验 + 传数据」，不碰计算")
arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.35), Inches(2.95), Inches(0.6), Inches(0.5))
arrow.fill.solid(); arrow.fill.fore_color.rgb = ACCENT; arrow.line.fill.background(); arrow.shadow.inherit = False
layer(3.55, "受约束编排层", "白名单 / 守卫 / grounding / 溯源（run_id+哈希） / 物理一致性检查")
arrow2 = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.35), Inches(4.75), Inches(0.6), Inches(0.5))
arrow2.fill.solid(); arrow2.fill.fore_color.rgb = ACCENT; arrow2.line.fill.background(); arrow2.shadow.inherit = False
layer(5.35, "领域工具链 phyxexp", "不确定度传播 / 线性拟合 / 伏安法·弹簧振子·杨氏模量 / 残差图·拟合图（JJF 1059.1-2012）")

# ============ Slide 5 核心机制 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "五道约束，让结论可信")
items = [
    ("白名单工具链", "计算能力只能来自 TOOL_REGISTRY（公式可查、可复现）"),
    ("守卫", "白名单外工具 / 任意代码执行 / 自行算数 → 直接拒绝"),
    ("grounding", "结论里的每个数字必须溯源到工具输出，否则判失败"),
    ("溯源 + 可复现", "每次运行产出 run_id + 输入哈希 + recipe，重放逐位一致"),
    ("物理一致性", "截距近零、R>0、R² 合理、残差无趋势，不通过则如实报告"),
]
for i, (h, b) in enumerate(items):
    yy = 1.65 + i * 1.12
    num = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.7), Inches(yy), Inches(0.62), Inches(0.62))
    num.fill.solid(); num.fill.fore_color.rgb = ACCENT; num.line.fill.background(); num.shadow.inherit = False
    ntf = num.text_frame; ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
    np = ntf.paragraphs[0]; np.alignment = PP_ALIGN.CENTER
    nr = np.add_run(); nr.text = str(i+1); style(nr, 18, True, WHITE)
    add_text(s, 1.6, yy, 11.0, 1.0, [(h, 18, True, PRIMARY)])
    add_text(s, 1.6, yy+0.42, 11.0, 0.6, [(b, 13, False, GRAY)])

# ============ Slide 6 实际使用场景 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "实际使用场景：一套工具链，三种宿主")
add_card(s, 0.6, 1.7, 3.95, 3.5, "Claude Code", [
    "加载根目录 SKILL.md",
    "",
    "大模型只被允许调用",
    "scripts/run.py --experiment <id>",
    "--data '<json>'",
    "",
    "数值全部来自工具链输出。",
], PRIMARY)
add_card(s, 4.85, 1.7, 3.95, 3.5, "ChatGPT（Custom GPT）", [
    "Instructions + Actions：",
    "挂 chatgpt/openapi.json",
    "",
    "大模型只调 HTTP POST /run，",
    "拿回结构化 JSON。",
    "",
    "用 Actions 而非 Code Interpreter，",
    "因为 Actions 天生是受约束调用。",
], GREEN)
add_card(s, 9.1, 1.7, 3.95, 3.5, "命令行 / 脚本", [
    "python scripts/run.py --list",
    "python scripts/run.py -e ohms_law ...",
    "",
    "python demo/compare.py",
    "（裸LLM vs 受约束Agent 对比）",
    "",
    "零依赖 HTTP 服务：",
    "python scripts/server.py",
], BLUE2)
add_text(s, 0.6, 5.6, 12.2, 1.0, [
    ("同一套 phyxexp 工具链与 config/experiments.json，三种宿主各自加载对应的「说明书」，约束始终在运行时这一层生效。", 15, True, PRIMARY),
])

# ============ Slide 7 效果对比 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "效果对比：裸 LLM vs 受约束 Agent", "伏安法测电阻 · 真实运行结果")
rows = [
    ("", "免费 AI", "phyxexp（受约束）"),
    ("电阻 R", "9.848571428571425 Ω（15位小数）", "9.85 ± 0.04 Ω（k=2，规范修约）"),
    ("不确定度", "± 0.5 Ω（拍脑袋）", "± 0.02 Ω（偏导传播）"),
    ("r / R²", "无", "0.99999 / 0.99998"),
    ("残差图", "无", "有（随机、无趋势）"),
    ("截距检查", "无", "b=0.005V，占比 0.09%，通过"),
    ("可复现", "否", "是（run_id + 重放一致）"),
]
tbl_shape = s.shapes.add_table(len(rows), 3, Inches(0.7), Inches(1.7), Inches(12.0), Inches(4.4))
tbl = tbl_shape.table
tbl.columns[0].width = Inches(2.2)
tbl.columns[1].width = Inches(4.4)
tbl.columns[2].width = Inches(5.4)
for ri, row in enumerate(rows):
    for ci, val in enumerate(row):
        cell = tbl.cell(ri, ci)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.08)
        tf = cell.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER if ci > 0 else PP_ALIGN.LEFT
        r = p.add_run(); r.text = val
        if ri == 0:
            style(r, 15, True, WHITE)
            cell.fill.solid(); cell.fill.fore_color.rgb = PRIMARY
        else:
            style(r, 14, ci == 2, GRAY if ci == 0 else (RED if ci == 1 else GREEN))
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if ci != 2 else RGBColor(0xEE, 0xF6, 0xEE)
add_text(s, 0.7, 6.35, 12.0, 0.6, [("一句话：裸 LLM 只有点估计；受约束 Agent 给 值±不确定度 + R² + 残差图 + 可复现证据。", 15, True, PRIMARY)])

# ============ Slide 8 演进路径 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "演进路径：三层递进")
steps = [
    ("短期", "物理实验", "今晚 Demo：线性拟合闭环（伏安法 / 弹簧振子 / 杨氏模量），跑通约束全流程。", BLUE2),
    ("中期", "跨学科实验工具包", "化学、生物等学科——加实验 = 在 config 加一条 JSON +（可选）登记一个公式工具。", GREEN),
    ("远期", "通用数据分析 Skill", "降级为通用数据分析 Agent / 宿主大模型的 Skill 插件，应用到工作与生活。", ACCENT),
]
x = 0.6
for tag, head, body, color in steps:
    add_card(s, x, 1.7, 4.0, 4.2, tag + " · " + head, [body], color)
    x += 4.15

# ============ Slide 9 开源技术路线 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "开源技术路线")
add_card(s, 0.7, 1.7, 3.8, 3.4, "phyxexp（MIT）", [
    "自研受约束引擎 + Agent 编排层，",
    "MIT 开源，代码可审计。",
    "",
    "仓库：GitHub（准备上传）",
], PRIMARY)
add_card(s, 5.0, 1.7, 3.8, 3.4, "依据 JJF 1059.1-2012", [
    "A 类 s/√n、B 类 Δ/√3、",
    "方和根合成——不确定度评定",
    "有国家标准背书。",
], GREEN)
add_card(s, 9.3, 1.7, 3.8, 3.4, "零重型依赖", [
    "核心引擎只依赖 numpy；",
    "绘图 matplotlib；",
    "HTTP 服务用标准库，零安装。",
    "",
    "可复现、可离线、可审计。",
], BLUE2)
add_text(s, 0.7, 5.5, 12.0, 1.0, [
    ("已覆盖 11 个本科物理实验（力/电/光/热），加实验 = 加配置；开源可逐行核对「约束到底在哪」。", 15, True, ACCENT),
])

# ============ Slide 10 落地价值 + 结尾 ============
s = prs.slides.add_slide(BLANK)
add_title(s, "落地价值")
vals = [
    ("学生", "实验报告不再「编误差」，结论带不确定度与残差图，可复现。"),
    ("教师 / 评测", "每个结果可溯源、可审计，能区分「真算」与「瞎编」。"),
    ("行业", "任何需要「可信计算」的场景：质检、工程、科研辅助。"),
]
y = 1.7
for h, b in vals:
    add_text(s, 0.9, y, 11.5, 1.1, [(h, 20, True, PRIMARY)])
    add_text(s, 0.9, y+0.5, 11.5, 0.7, [(b, 15, False, GRAY)])
    y += 1.5
end = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.6), Inches(6.1), Inches(10.1), Inches(1.0))
end.fill.solid(); end.fill.fore_color.rgb = PRIMARY; end.line.fill.background(); end.shadow.inherit = False
etf = end.text_frame; etf.vertical_anchor = MSO_ANCHOR.MIDDLE
ep = etf.paragraphs[0]; ep.alignment = PP_ALIGN.CENTER
er = ep.add_run(); er.text = "phyxexp —— 约束，本身就是可信度的来源。"
style(er, 22, True, WHITE)

out = "docs/phyxexp-参赛PPT-初版.pptx"
prs.save(out)
print("saved", out, "slides:", len(prs.slides._sldIdLst))
